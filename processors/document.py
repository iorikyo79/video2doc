"""
DocumentProcessor 클래스 - 문서 변환 관련 기능을 담당
"""

import logging
from pathlib import Path
from typing import Optional

from markitdown import MarkItDown
from PIL import Image
from pptx import Presentation
from utils.error_handlers import handle_error_with_suggestions

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """문서 파일 변환을 담당하는 클래스"""
    
    def __init__(self, output_dir: str = "output", markitdown_instance: Optional[MarkItDown] = None):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)
        self.markitdown = markitdown_instance
    
    def convert_pptx_to_md(self, pptx_path: str) -> str:
        """PPTX를 마크다운으로 변환 (레거시 방식 - markitdown 실패 시 사용)"""
        pptx_file = Path(pptx_path)
        md_file = self.output_dir / f"04_{pptx_file.stem}_ref.md"
        
        try:
            prs = Presentation(pptx_path)
            content = f"# {pptx_file.name} 내용\n\n"
            
            for slide_num, slide in enumerate(prs.slides, 1):
                content += f"## 슬라이드 {slide_num}\n\n"
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        content += f"{shape.text}\n\n"
            
            with open(md_file, 'w', encoding='utf-8') as f:
                f.write(content)
            
            logger.info(f"PPTX 변환 완료: {md_file}")
            return str(md_file)
            
        except Exception as e:
            logger.error(f"PPTX 변환 실패: {e}")
            raise handle_error_with_suggestions(e, "document", pptx_path)
    
    def convert_image_to_md(self, image_path: str) -> str:
        """이미지를 마크다운으로 변환 (markitdown 실패 시 기본 정보만 제공)"""
        image_file = Path(image_path)
        md_file = self.output_dir / f"04_{image_file.stem}_ref.md"
        
        try:
            # 이미지 기본 정보 제공
            image = Image.open(image_path)
            width, height = image.size
            
            content = f"# {image_file.name} 이미지 정보\n\n"
            content += f"**이미지 파일:** {image_file.name}\n"
            content += f"**파일 형식:** {image_file.suffix.upper()}\n"
            content += f"**크기:** {width} x {height} pixels\n\n"
            content += f"## 참고\n\n"
            content += f"이미지 내용 분석을 위해서는 markitdown의 LLM 연동 기능을 사용하세요.\n\n"
            
            with open(md_file, 'w', encoding='utf-8') as f:
                f.write(content)
            
            logger.info(f"이미지 정보 변환 완료: {md_file}")
            return str(md_file)
            
        except Exception as e:
            logger.error(f"이미지 정보 변환 실패: {e}")
            raise handle_error_with_suggestions(e, "document", image_path)

    def convert_reference_file_to_md(self, file_path: str) -> str:
        """
        참조 파일을 마크다운으로 변환 (markitdown 통합 방식)
        
        Args:
            file_path: 변환할 파일 경로
            
        Returns:
            변환된 마크다운 파일 경로
        """
        input_file = Path(file_path)
        md_file = self.output_dir / f"04_{input_file.stem}_ref.md"
        
        try:
            if self.markitdown is None:
                logger.warning("markitdown이 초기화되지 않음, 레거시 방식으로 변환 시도")
                return self._convert_file_legacy(file_path)
            
            logger.info(f"markitdown으로 파일 변환 시작: {file_path}")
            
            # 이미지 파일인 경우 특별 처리 옵션 확인
            if input_file.suffix.lower() in ['.jpg', '.jpeg', '.png', '.bmp', '.gif', '.tiff', '.webp']:
                return self._convert_image_with_markitdown(file_path)
            
            # markitdown을 사용한 일반 문서 변환
            result = self.markitdown.convert(file_path)
            
            # 변환 결과를 마크다운 형식으로 저장
            content = f"# {input_file.name} 내용\n\n"
            content += f"**파일 형식:** {input_file.suffix.upper()}\n"
            content += f"**변환 도구:** markitdown (Microsoft)\n"
            content += f"**원본 파일:** {input_file.name}\n\n"
            content += "---\n\n"
            content += result.text_content
            
            with open(md_file, 'w', encoding='utf-8') as f:
                f.write(content)
            
            logger.info(f"markitdown 변환 완료: {md_file}")
            return str(md_file)
            
        except Exception as e:
            logger.warning(f"markitdown 변환 실패: {e}, 레거시 방식으로 시도")
            try:
                return self._convert_file_legacy(file_path)
            except Exception as legacy_error:
                logger.error(f"레거시 변환도 실패: {legacy_error}")
                raise
    
    def _convert_image_with_markitdown(self, image_path: str) -> str:
        """
        이미지 파일을 markitdown을 사용하여 변환 (LLM 지원 고려)
        
        Args:
            image_path: 이미지 파일 경로
            
        Returns:
            변환된 마크다운 파일 경로
        """
        input_file = Path(image_path)
        md_file = self.output_dir / f"04_{input_file.stem}_ref.md"
        
        try:
            # markitdown으로 이미지 처리 시도 (LLM 연동 시 더 나은 결과)
            result = self.markitdown.convert(image_path)
            
            content = f"# {input_file.name} 이미지 분석\n\n"
            content += f"**파일 형식:** {input_file.suffix.upper()}\n"
            content += f"**변환 도구:** markitdown (Microsoft)\n"
            content += f"**원본 파일:** {input_file.name}\n\n"
            content += "## 이미지 내용 분석\n\n"
            content += result.text_content
            
            with open(md_file, 'w', encoding='utf-8') as f:
                f.write(content)
                
            logger.info(f"markitdown 이미지 변환 완료: {md_file}")
            return str(md_file)
            
        except Exception as e:
            logger.warning(f"markitdown 이미지 변환 실패: {e}, 기본 이미지 정보로 대체")
            return self.convert_image_to_md(image_path)
    
    def _convert_file_legacy(self, file_path: str) -> str:
        """
        레거시 방식으로 파일 변환 (fallback)
        
        Args:
            file_path: 변환할 파일 경로
            
        Returns:
            변환된 마크다운 파일 경로
        """
        file_path_obj = Path(file_path)
        suffix = file_path_obj.suffix.lower()
        
        if suffix == '.pptx':
            return self.convert_pptx_to_md(file_path)
        elif suffix in ['.jpg', '.jpeg', '.png', '.bmp', '.gif', '.tiff', '.webp']:
            return self.convert_image_to_md(file_path)
        elif suffix in ['.txt', '.md']:
            # 텍스트 파일은 단순 복사
            return self._convert_text_file(file_path)
        else:
            raise ValueError(f"지원하지 않는 파일 형식: {suffix}")
    
    def _convert_text_file(self, text_path: str) -> str:
        """
        텍스트 파일을 마크다운으로 변환 (단순 복사)
        
        Args:
            text_path: 텍스트 파일 경로
            
        Returns:
            변환된 마크다운 파일 경로
        """
        text_file = Path(text_path)
        md_file = self.output_dir / f"04_{text_file.stem}_ref.md"
        
        try:
            with open(text_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # 마크다운 헤더 추가
            md_content = f"# {text_file.name} 내용\n\n"
            md_content += f"**파일 형식:** {text_file.suffix.upper()}\n"
            md_content += f"**원본 파일:** {text_file.name}\n\n"
            md_content += "---\n\n"
            md_content += content
            
            with open(md_file, 'w', encoding='utf-8') as f:
                f.write(md_content)
            
            logger.info(f"텍스트 파일 변환 완료: {md_file}")
            return str(md_file)
            
        except Exception as e:
            logger.error(f"텍스트 파일 변환 실패: {e}")
            raise handle_error_with_suggestions(e, "document", text_path)
